"""Extract selected IMF CDIS tables from project-local, checksummed workbooks."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pandas as pd

from gtri_config import COUNTRIES, OUTPUT_DIR, PROCESSED_DIR, PROJECT_ROOT, RAW_DIR, append_progress, ensure_directories


YEARS = (2017, 2021)
INPUTS = tuple(
    (year, flow, RAW_DIR / "fdi" / f"{year}_{flow}.xlsx")
    for year in YEARS
    for flow in ("inward", "outward")
)
COUNTRY_NAMES = {country["country"] for country in COUNTRIES}


def parse_value(value: object) -> float | None:
    if value in (None, "C", "..", "—", ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def extract_workbook(year: int, flow: str, path: Path) -> tuple[list[dict], dict]:
    if not path.is_file():
        raise FileNotFoundError(f"Required supplied IMF CDIS workbook is missing: {path}")

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    # These supplied sheets are modest in size. Normal mode keeps repeated
    # cell lookups in memory; read_only mode reparses sheet XML on random access.
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook.active
        title = str(sheet.cell(2, 1).value or "")
        expected_marker = "Table 6-i" if flow == "inward" else "Table 6-o"
        if expected_marker not in title or str(year) not in title:
            raise RuntimeError(f"Unexpected CDIS table title in {path.name}: {title}")

        headers = [sheet.cell(4, col).value for col in range(1, sheet.max_column + 1)]
        columns = {str(value).strip(): index + 1 for index, value in enumerate(headers) if value is not None}
        china_col = columns.get("China, P.R.: Mainland")
        china_row = next(
            (row for row in range(5, sheet.max_row + 1)
             if str(sheet.cell(row, 1).value or "").strip() == "China, P.R.: Mainland"),
            None,
        )

        rows: list[dict] = []
        for economy in COUNTRIES:
            if flow == "inward":
                row = next(
                    (r for r in range(5, sheet.max_row + 1)
                     if str(sheet.cell(r, 1).value or "").strip() == economy["country"]),
                    None,
                )
                raw_value = sheet.cell(row, china_col).value if row is not None and china_col else None
            else:
                country_col = columns.get(economy["country"])
                raw_value = sheet.cell(china_row, country_col).value if china_row is not None and country_col else None

            flow_name = "china_inward_fdi_usd_millions" if flow == "inward" else "china_outward_fdi_usd_millions"
            rows.append(
                {
                    "year": year,
                    "country": economy["country"],
                    "iso3": economy["iso3"],
                    "flow": flow_name,
                    "value_usd_millions": parse_value(raw_value),
                    "source_file": path.name,
                    "source_artifact": path.relative_to(PROJECT_ROOT).as_posix(),
                    "source_sha256": digest,
                    "source_flag": "IMF_CDIS_reporter_positions_by_immediate_counterpart",
                }
            )

        manifest = {
            "file": path.name,
            "year": year,
            "table_type": flow,
            "title": title,
            "source_artifact": path.relative_to(PROJECT_ROOT).as_posix(),
            "bytes": path.stat().st_size,
            "sha256": digest,
            "rows_extracted": len(rows),
        }
        return rows, manifest
    finally:
        workbook.close()


def main() -> int:
    ensure_directories()
    rows: list[dict] = []
    manifests: list[dict] = []
    for year, flow, path in INPUTS:
        extracted, manifest = extract_workbook(year, flow, path)
        rows.extend(extracted)
        manifests.append(manifest)

    output = pd.DataFrame(rows)
    keys = ["year", "iso3", "flow"]
    if len(output) != 4 * len(COUNTRIES) or output.duplicated(keys).any():
        raise RuntimeError("CDIS input does not produce exactly one row per economy-year-flow key")
    expected = {
        (year, economy["iso3"], flow)
        for year in YEARS
        for economy in COUNTRIES
        for flow in ("china_inward_fdi_usd_millions", "china_outward_fdi_usd_millions")
    }
    if set(map(tuple, output[keys].itertuples(index=False, name=None))) != expected:
        raise RuntimeError("CDIS coverage does not match the expected 12 economies x 2 years x 2 flows")

    output_path = PROCESSED_DIR / "cdis_china_fdi_key_years.csv"
    output.to_csv(output_path, index=False, encoding="utf-8-sig")
    manifest = pd.DataFrame(manifests)
    manifest_path = RAW_DIR / "fdi" / "key_year_manifest.csv"
    manifest.to_csv(manifest_path, index=False, encoding="utf-8-sig")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output.to_csv(OUTPUT_DIR / output_path.name, index=False, encoding="utf-8-sig")
    manifest.to_csv(OUTPUT_DIR / manifest_path.name, index=False, encoding="utf-8-sig")

    missing = int(output["value_usd_millions"].isna().sum())
    negative = int(output["value_usd_millions"].lt(0).sum())
    append_progress(
        "Supplied IMF CDIS key-year FDI module",
        "complete as optional non-scoring context",
        [
            f"Read {len(manifest)} project-local workbooks from data/raw/fdi and extracted {len(output)} unique economy-year-flow rows.",
            "Retained inward and outward positions separately; values are end-year positions in USD millions by immediate counterpart, not annual flows.",
            "Selected target-year links are 2017 to 2017 (lag 0) and 2022 to 2021 (lag 1); CDIS has no supplied 2022 workbook and is not semiconductor-specific.",
            f"Observed missing/nonreported cells: {missing}; negative net position cells: {negative}; no imputation and no machine-specific file paths retained.",
        ],
    )
    print(f"CDIS rows={len(output)}; files={len(manifest)}; missing={missing}; negative={negative}")
    print(manifest[["file", "table_type", "sha256"]].to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
